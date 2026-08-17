import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule

# ----------------------------------------------------
# File Paths
# ----------------------------------------------------

FINANCIAL_RATIOS = "output/financial_ratios.csv"
PEER_PERCENTILES = "output/peer_percentiles.csv"
PEER_GROUPS = "data\\raw\\peer_groups.xlsx"

# Change this if you exported companies.csv instead
COMPANIES = "data\\raw\\companies.xlsx"

OUTPUT_FILE = "output/peer_comparison.xlsx"

# ----------------------------------------------------
# Auto Column Width
# ----------------------------------------------------

def auto_adjust_columns(ws):

    for column in ws.columns:

        max_length = 0
        letter = column[0].column_letter

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[letter].width = max_length + 3


# ----------------------------------------------------
# Conditional Formatting
# ----------------------------------------------------

def add_conditional_formatting(ws):

    if ws.max_row <= 1:
        return

    rule = ColorScaleRule(

        start_type="min",
        start_color="F8696B",

        mid_type="percentile",
        mid_value=50,
        mid_color="FFEB84",

        end_type="max",
        end_color="63BE7B",
    )

    for col in range(4, ws.max_column + 1):

        letter = ws.cell(
            row=1,
            column=col
        ).column_letter

        ws.conditional_formatting.add(

            f"{letter}2:{letter}{ws.max_row}",
            rule,
        )


# ----------------------------------------------------
# Main
# ----------------------------------------------------

def generate_peer_comparison():

    print("=" * 70)
    print("Generating Peer Comparison Workbook")
    print("=" * 70)

    # ----------------------------------------
    # Load Files
    # ----------------------------------------

    financial = pd.read_csv(FINANCIAL_RATIOS)

    peer_percentiles = pd.read_csv(PEER_PERCENTILES)

    peer_groups = pd.read_excel(PEER_GROUPS)

    companies = pd.read_excel(COMPANIES,header=1)
    

    print("Financial Ratios :", len(financial))
    print("Peer Percentiles :", len(peer_percentiles))
    print("Peer Groups :", len(peer_groups))
    print("Companies :", len(companies))

    # ----------------------------------------
    # Merge
    # ----------------------------------------

    df = financial.merge(

        peer_groups[
            [
                "company_id",
                "peer_group_name",
            ]
        ],

        on="company_id",

        how="left",

    )

    df = df.merge(

        companies[
            [
                "id",
                "company_name",
            ]
        ],

        left_on="company_id",

        right_on="id",

        how="left",

    )

    # ----------------------------------------
    # Metrics
    # ----------------------------------------

    metrics = [

        "return_on_equity_pct",

        "roce_calculated",

        "operating_profit_margin_pct",

        "net_profit_margin_pct",

        "debt_to_equity",

        "interest_coverage",

        "asset_turnover",

        "free_cash_flow",

        "cfo_quality_score",

        "capex_intensity",

        "fcf_conversion",

    ]

    # ----------------------------------------
    # Create Workbook
    # ----------------------------------------

    with pd.ExcelWriter(

        OUTPUT_FILE,

        engine="openpyxl"

    ) as writer:

        groups = sorted(

            df["peer_group_name"]

            .dropna()

            .unique()

        )

        print()

        print("Peer Groups :", len(groups))

        print()

        for group in groups:

            print("Creating Sheet :", group)

            temp = df[

                df["peer_group_name"] == group

            ].copy()

            if "year" in temp.columns:

                temp = (

                    temp.sort_values("year")

                    .drop_duplicates(

                        "company_id",

                        keep="last",

                    )

                )

            columns = [

                "company_name",

                "company_id",

                "year",

            ]

            for metric in metrics:

                if metric in temp.columns:

                    columns.append(metric)

            temp = temp[columns]

            sheet_name = str(group)[:31]

            temp.to_excel(

                writer,

                sheet_name=sheet_name,

                index=False,

            )

    # ----------------------------------------
    # Formatting
    # ----------------------------------------

    workbook = load_workbook(OUTPUT_FILE)

    for sheet in workbook.sheetnames:

        ws = workbook[sheet]

        auto_adjust_columns(ws)

        add_conditional_formatting(ws)

    workbook.save(OUTPUT_FILE)

    print()

    print("=" * 70)
    print("peer_comparison.xlsx generated successfully!")
    print("=" * 70)
    print("Saved :", OUTPUT_FILE)


# ----------------------------------------------------

if __name__ == "__main__":

    generate_peer_comparison()